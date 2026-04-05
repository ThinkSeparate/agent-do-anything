# tools/office_tools.py
from langchain_core.tools import tool
from docx import Document
from pptx import Presentation
from pptx.util import Inches


@tool
def create_ppt(file_path: str, title: str = "Presentation") -> str:
    """创建一个新的PPT文件，并添加一个标题幻灯片。"""
    try:
        prs = Presentation()
        slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(slide_layout)
        title_shape = slide.shapes.title
        title_shape.text = title
        prs.save(file_path)
        return f"PPT文件已创建于 '{file_path}'，标题为 '{title}'。"
    except Exception as e:
        return f"创建PPT文件时出错: {e}"


@tool
def add_slide_to_ppt(file_path: str, title: str = "", content: str = "") -> str:
    """向现有PPT文件添加一个幻灯片。"""
    try:
        prs = Presentation(file_path)
        slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(slide_layout)
        title_shape = slide.shapes.title
        title_shape.text = title
        content_shape = slide.placeholders[1]
        content_shape.text = content
        prs.save(file_path)
        return f"已向PPT文件添加幻灯片，标题: '{title}'。"
    except FileNotFoundError:
        return f"错误：找不到文件 '{file_path}'。"
    except Exception as e:
        return f"添加幻灯片时出错: {e}"


@tool
def read_ppt(file_path: str) -> str:
    """读取PPT文件并返回幻灯片标题和内容的文本表示。"""
    try:
        prs = Presentation(file_path)
        slides_data = []
        for i, slide in enumerate(prs.slides):
            title = slide.shapes.title.text if slide.shapes.title else ""
            content = ""
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape != slide.shapes.title:
                    content += shape.text + "\n"
            slides_data.append(f"幻灯片 {i}: 标题='{title}', 内容='{content.strip()}'")
        return "\n".join(slides_data)
    except FileNotFoundError:
        return f"错误：找不到文件 '{file_path}'。"
    except Exception as e:
        return f"读取PPT文件时出错: {e}"


@tool
def modify_ppt_slide(file_path: str, slide_index: int, new_title: str = None, new_content: str = None) -> str:
    """修改指定索引的幻灯片。"""
    try:
        prs = Presentation(file_path)
        if 0 <= slide_index < len(prs.slides):
            slide = prs.slides[slide_index]
            if new_title is not None and slide.shapes.title:
                slide.shapes.title.text = new_title
            if new_content is not None:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape != slide.shapes.title:
                        shape.text = new_content
                        break
            prs.save(file_path)
            return f"已成功更新第 {slide_index} 张幻灯片。"
        else:
            return f"错误：幻灯片索引 {slide_index} 超出范围。"
    except FileNotFoundError:
        return f"错误：找不到文件 '{file_path}'。"
    except Exception as e:
        return f"修改幻灯片时出错: {e}"


@tool
def read_docx(file_path: str) -> str:
    """读取 DOCX 文件并返回其纯文本内容。"""
    try:
        doc = Document(file_path)
        full_text = []
        for paragraph in doc.paragraphs:
            full_text.append(paragraph.text)
        return "\n".join(full_text)
    except Exception as e:
        return f"读取DOCX文件时出错: {e}"


@tool
def modify_docx_paragraph(file_path: str, paragraph_index: int, new_text: str) -> str:
    """修改 DOCX 文件中指定段落的文本。"""
    try:
        doc = Document(file_path)
        if 0 <= paragraph_index < len(doc.paragraphs):
            doc.paragraphs[paragraph_index].text = new_text
            doc.save(file_path)
            return f"已成功更新第 {paragraph_index} 段内容。"
        else:
            return f"错误：段落索引 {paragraph_index} 超出范围。"
    except Exception as e:
        return f"修改DOCX段落时出错: {e}"


office_tools = [create_ppt, add_slide_to_ppt, read_ppt, modify_ppt_slide, read_docx, modify_docx_paragraph]
